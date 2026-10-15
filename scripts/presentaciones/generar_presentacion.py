"""Genera la presentación de cierre de ciclo de un agricultor a partir de la plantilla.

Uso:
    python scripts/presentaciones/extraer_datos.py B01 > datos_B01.json
    python scripts/presentaciones/generar_presentacion.py datos_B01.json salida.pptx \
        [--plantilla docs/presentaciones/plantilla_agricultores.pptx] [--nombre-corto "JULIO COLMENAREZ"]

La plantilla es la presentación de Lorena Orlando. Las láminas comunes (portada,
plan agronómico, formaciones, I+D, agricultura regenerativa, I+D 2027) se dejan
tal cual; las que dependen del agricultor se reescriben clonando las formas de
la propia plantilla, para que tipografía y colores sigan siendo los suyos.

Lo que no está en la base (pagos de mecanización y flujo de caja) se escribe
como PENDIENTE: búscalo con `markitdown salida.pptx | grep PENDIENTE`.
"""

import argparse
import copy
import io
import json
import math
import os
import re
import urllib.request
from datetime import date

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Inches, Pt

AQUI = os.path.dirname(os.path.abspath(__file__))
PLANTILLA = os.path.join(AQUI, "..", "..", "docs", "presentaciones", "plantilla_agricultores.pptx")
PENDIENTE = "PENDIENTE"
ESTADO_INCIDENCIA = "Atendida"  # el tag de la plantilla de Lorena

VERDE_OSC, VERDE, VERDE_CLARO, CREMA = "1F2E18", "4C6A35", "88A85E", "EDF0E2"
DORADO, OXIDO, AZUL, GRIS = "C9962E", "B4533A", "34506B", "6B7560"
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
MESES = ["", "ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO", "AGOSTO",
         "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]


# ------------------------------------------------------------------ formato
def es(n, dec=2):
    """Número con formato venezolano: 1.234,56"""
    s = f"{n:,.{dec}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def ddmm(iso):
    d = date.fromisoformat(iso)
    return f"{d.day:02d}/{d.month:02d}"


def nombre_corto(nombre):
    t = nombre.split()
    if len(t) >= 4:
        return f"{t[0]} {t[2]}"
    return f"{t[0]} {t[-1]}" if len(t) > 1 else nombre


# ------------------------------------------------------------------ formas
def forma(slide, nombre, n=0):
    """La n-ésima forma con ese nombre (la plantilla repite algunos nombres)."""
    hits = []

    def rec(shapes):
        for s in shapes:
            if s.name == nombre:
                hits.append(s)
            if s.shape_type == 6:
                rec(s.shapes)
    rec(slide.shapes)
    return hits[n]


def por_id(slide, shape_id):
    def rec(shapes):
        for s in shapes:
            if s.shape_id == shape_id:
                return s
            if s.shape_type == 6:
                r = rec(s.shapes)
                if r is not None:
                    return r
    s = rec(slide.shapes)
    if s is None:
        raise KeyError(shape_id)
    return s


def borrar(*shapes):
    for s in shapes:
        s._element.getparent().remove(s._element)


def _nuevo_id(slide):
    ids = [int(x) for x in slide._element.xpath(".//p:cNvPr/@id")]
    return max(ids) + 1


def clonar(slide, s, x=None, y=None, w=None, h=None):
    el = copy.deepcopy(s._element)
    nv = el.xpath(".//p:cNvPr")[0]
    nv.set("id", str(_nuevo_id(slide)))
    nv.set("name", nv.get("name") + " (gen)")
    for ext in el.xpath(".//p:cNvPr/a:extLst"):
        ext.getparent().remove(ext)
    slide.shapes._spTree.append(el)
    nueva = next(sh for sh in slide.shapes if sh._element is el)
    mover(nueva, x, y, w, h)
    return nueva


def mover(s, x=None, y=None, w=None, h=None):
    if x is not None:
        s.left = Inches(x)
    if y is not None:
        s.top = Inches(y)
    if w is not None:
        s.width = Inches(w)
    if h is not None:
        s.height = Inches(h)


def texto(s, lineas, like=None, sz=None, color=None, bold=None, align=None):
    """Reemplaza el texto conservando el formato del primer run (o el de `like`).

    `lineas` puede ser un str (con \\n) o una lista. Cada línea puede ser una
    lista de runs [(texto, {"color":..,"sz":..,"b":..}), ...].
    """
    if isinstance(lineas, str):
        lineas = lineas.split("\n")
    origen = (like or s).text_frame._txBody
    p0 = origen.xpath("./a:p")[0]
    r0 = (origen.xpath(".//a:r") or [None])[0]
    body = s.text_frame._txBody
    for p in body.xpath("./a:p"):
        body.remove(p)
    for linea in lineas:
        p = copy.deepcopy(p0)
        for hijo in list(p):
            if hijo.tag != f"{{{NS['a']}}}pPr":
                p.remove(hijo)
        runs = linea if isinstance(linea, list) else [(linea, {})]
        for txt, opt in runs:
            if r0 is not None:
                r = copy.deepcopy(r0)
            else:
                r = etree.Element(f"{{{NS['a']}}}r")
                etree.SubElement(r, f"{{{NS['a']}}}rPr", lang="es-VE")
                etree.SubElement(r, f"{{{NS['a']}}}t")
            p.append(r)
            rpr = r.find("a:rPr", NS)
            r.find("a:t", NS).text = txt
            _estilo_run(rpr, sz=opt.get("sz", sz), color=opt.get("color", color), bold=opt.get("b", bold))
        if align:
            ppr = p.find("a:pPr", NS)
            if ppr is not None:
                ppr.set("algn", align)
        body.append(p)
    return s


def _estilo_run(rpr, sz=None, color=None, bold=None):
    if sz is not None:
        rpr.set("sz", str(int(sz * 100)))
    if bold is not None:
        rpr.set("b", "1" if bold else "0")
    if color is not None:
        for f in rpr.findall("a:solidFill", NS):
            rpr.remove(f)
        fill = etree.Element(f"{{{NS['a']}}}solidFill")
        etree.SubElement(fill, f"{{{NS['a']}}}srgbClr").set("val", color)
        rpr.insert(0, fill)


def runs(s, textos):
    """Cambia el texto de cada run en orden, sin tocar su formato."""
    rs = s.text_frame._txBody.xpath(".//a:r")
    for r, t in zip(rs, textos):
        r.find("a:t", NS).text = t
    for r in rs[len(textos):]:
        r.getparent().remove(r)


def relleno(s, color):
    s.fill.solid()
    s.fill.fore_color.rgb = RGBColor.from_string(color)


def tarjeta(slide, x, y, w, h, fondo="FFFFFF", borde="DCE2CC"):
    r = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    r.adjustments[0] = 0.12
    relleno(r, fondo)
    if borde:
        r.line.color.rgb = RGBColor.from_string(borde)
        r.line.width = Pt(0.75)
    else:
        r.line.fill.background()
    r.shadow.inherit = False
    return r


def reemplazar_imagen(slide, pic, blob):
    """Cambia la imagen de un Picture conservando posición, recorte y z-order."""
    img_part, rId = pic.part.get_or_add_image_part(io.BytesIO(blob))
    pic._element.xpath(".//a:blip")[0].set(
        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed", rId)
    pic.crop_left = pic.crop_right = pic.crop_top = pic.crop_bottom = 0


# ------------------------------------------------------------------ imágenes
def satelital(lat, lng, w_px=1400, h_px=700, zoom=15):
    """Mosaico de Esri World Imagery centrado en la finca."""
    from PIL import Image, ImageDraw
    n = 2 ** zoom
    xt = (lng + 180) / 360 * n
    yt = (1 - math.asinh(math.tan(math.radians(lat))) / math.pi) / 2 * n
    cx, cy = xt * 256, yt * 256
    x0, y0 = int(cx - w_px / 2), int(cy - h_px / 2)
    img = Image.new("RGB", (w_px, h_px))
    for tx in range(x0 // 256, (x0 + w_px) // 256 + 1):
        for ty in range(y0 // 256, (y0 + h_px) // 256 + 1):
            url = f"https://server.arcgisonline.com/ArcGIS/rest/services/World_Imagery/MapServer/tile/{zoom}/{ty}/{tx}"
            with urllib.request.urlopen(urllib.request.Request(url, headers={"User-Agent": "saturno"}), timeout=60) as r:
                tile = Image.open(io.BytesIO(r.read())).convert("RGB")
            img.paste(tile, (tx * 256 - x0, ty * 256 - y0))
    d = ImageDraw.Draw(img)
    px, py = w_px / 2, h_px / 2
    d.ellipse([px - 11, py - 11, px + 11, py + 11], outline="white", width=4)
    d.ellipse([px - 5, py - 5, px + 5, py + 5], fill="#C9962E")
    out = io.BytesIO()
    img.save(out, "JPEG", quality=88)
    return out.getvalue()


def grafico_lluvia(datos):
    """Curva mensual de lluvia: real (estación) + estimada (pronóstico), siembra y floración."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    import numpy as np

    fam = "Inter" if any("Inter" in f.name for f in font_manager.fontManager.ttflist) else "DejaVu Sans"
    plt.rcParams.update({"font.family": fam, "font.size": 11})
    c = datos["clima"]
    real = {m["mes"]: m["mm"] for m in c["mensual"] if m["cobertura"] >= 0.35}
    parcial = {m["mes"] for m in c["mensual"] if m["cobertura"] < 0.9}
    # Un mes con poca cobertura y pronóstico disponible se grafica como estimado.
    pron = {p["mes"]: p["mm"] for p in c["pronostico"]}
    for mes in list(real):
        if mes in pron and mes in parcial and mes == max(real):
            del real[mes]
    m_real = sorted(real)
    m_est = [m for m in sorted(pron) if m > max(m_real)]
    m0 = min(m_real)

    def suave(xs, ys, k=24):
        xs, ys = np.array(xs, float), np.array(ys, float)
        if len(xs) < 3:
            return xs, ys
        out_x, out_y = [], []
        p = np.concatenate([[ys[0]], ys, [ys[-1]]])
        for i in range(len(xs) - 1):
            t = np.linspace(0, 1, k, endpoint=False)
            p0, p1, p2, p3 = p[i], p[i + 1], p[i + 2], p[i + 3]
            seg = 0.5 * ((2 * p1) + (-p0 + p2) * t + (2 * p0 - 5 * p1 + 4 * p2 - p3) * t ** 2 + (-p0 + 3 * p1 - 3 * p2 + p3) * t ** 3)
            out_x += list(xs[i] + t * (xs[i + 1] - xs[i]))
            out_y += list(np.maximum(seg, 0))
        return np.array(out_x + [xs[-1]]), np.array(out_y + [ys[-1]])

    fig, ax = plt.subplots(figsize=(12.73, 3.35), dpi=200)
    xr, yr = suave(m_real, [real[m] for m in m_real])
    ax.plot(xr, yr, color="#2E8B3E", lw=2.2)
    if m_est:
        xe, ye = suave([m_real[-1]] + m_est, [real[m_real[-1]]] + [pron[m] for m in m_est])
        ax.plot(xe, ye, color="#2E8B3E", lw=2.2, ls=(0, (4, 3)))
    for m in m_real:
        ax.plot(m, real[m], "o", ms=5, color="#2E8B3E", mfc="white" if m in parcial else "#2E8B3E")
        ax.annotate(f"{real[m]:.0f}", (m, real[m]), textcoords="offset points", xytext=(0, 8), ha="center",
                    fontsize=9, color="#1F2E18")
    for m in m_est:
        ax.annotate(f"{pron[m]:.0f}", (m, pron[m]), textcoords="offset points", xytext=(0, 8), ha="center",
                    fontsize=9, color="#6B7560")

    def frac(iso):
        d = date.fromisoformat(iso)
        return d.month + (d.day - 1) / 31

    siembras = sorted(l["fecha_siembra"] for l in datos["lotes"] if l["fecha_siembra"])
    ymax = max(list(real.values()) + list(pron.values())) * 1.25
    if siembras:
        s0, s1 = frac(siembras[0]), frac(siembras[-1])
        ax.axvspan(s0, max(s1, s0 + 0.05), color="#34506B", alpha=0.12, lw=0)
        ax.axvline((s0 + s1) / 2, color="#34506B", ls=(0, (3, 3)), lw=1.4)
    if c["floracion"]:
        f0 = frac(min(f["desde"] for f in c["floracion"]))
        f1 = frac(max(f["hasta"] for f in c["floracion"]))
        ax.add_patch(plt.Rectangle((f0, 0), f1 - f0, ymax * 0.97, fill=False, ec="#B4533A", lw=2.2, ls=(0, (5, 3))))
    ax.set_xlim(m0 - 0.15, 12.15)
    ax.set_ylim(0, ymax)
    ax.set_xticks(range(m0, 13))
    ax.set_xticklabels(["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"][m0 - 1:],
                       color="#6B7560")
    ax.set_ylabel("mm", color="#6B7560")
    ax.tick_params(axis="y", colors="#6B7560", labelsize=9)
    ax.grid(axis="both", color="#E5E7E0", ls=(0, (3, 3)), lw=0.8)
    for lado in ("top", "right"):
        ax.spines[lado].set_visible(False)
    for lado in ("left", "bottom"):
        ax.spines[lado].set_color("#A9B09C")
    fig.tight_layout()
    out = io.BytesIO()
    fig.savefig(out, format="png", transparent=False, facecolor="white")
    plt.close(fig)
    return out.getvalue(), parcial & set(m_real)


# ------------------------------------------------------------------ láminas
def lamina_portada_agricultor(sl, d, corto):
    texto(forma(sl, "TextBox 7"), corto.upper())
    texto(forma(sl, "TextBox 8"), d["estado"].upper())


MAPAS = os.path.join(AQUI, "..", "..", "docs", "presentaciones", "mapas")


def mapa_convenio(codigo_up):
    """Mapa del Anexo B del convenio (recortado a mano) y su tabla sub-lote → lote de Saturno."""
    img, meta = os.path.join(MAPAS, f"{codigo_up}.jpg"), os.path.join(MAPAS, f"{codigo_up}.json")
    if not (os.path.exists(img) and os.path.exists(meta)):
        return None
    return open(img, "rb").read(), json.load(open(meta, encoding="utf-8"))


def lamina_suelo_mapa(sl, d):
    lotes, suelo = d["lotes"], d["suelo"]
    ha = sum(l["ha"] for l in lotes)
    pic = forma(sl, "Picture 22")
    pill = forma(sl, "Rounded Rectangle 23")
    cab = forma(sl, "TextBox 39")
    cuerpo = forma(sl, "TextBox 57")
    filas = sorted(lotes, key=lambda l: -l["ha"])
    borrar(forma(sl, "Picture 24"), forma(sl, "Rounded Rectangle 25"))
    conv = mapa_convenio(d["codigo_up"])

    if conv:
        # Mapa del convenio a la izquierda; al lado, qué sub-lotes del mapa forman cada lote.
        blob, meta = conv
        from PIL import Image
        w_px, h_px = Image.open(io.BytesIO(blob)).size
        ancho = min(4.9, 5.2 * w_px / h_px)
        alto_mapa = ancho * h_px / w_px
        reemplazar_imagen(sl, pic, blob)
        mover(pic, 0.45, 0.75, ancho, alto_mapa)
        x0 = 0.45 + ancho + 0.15
        ancho_t = 6.75 - x0
        clonar(sl, forma(sl, "Rounded Rectangle 56"), x0, 0.75, ancho_t, alto_mapa)
        texto(clonar(sl, cab, x0 + 0.12, 0.88, ancho_t - 0.2, 0.22), "Nº EN EL MAPA", sz=7.5)
        paso = min(0.5, (alto_mapa - 0.75) / max(len(filas), 1))
        for i, l in enumerate(filas):
            y = 1.2 + i * paso
            texto(clonar(sl, cuerpo, x0 + 0.12, y, ancho_t - 0.2, paso),
                  [[(meta["ids"].get(l["nombre"], "—"), {"color": DORADO, "b": True, "sz": 9.5})],
                   [(l["nombre"], {"color": VERDE_OSC, "b": True, "sz": 9})]])
        # Etiqueta de la finca bajo el mapa, como cinta.
        runs(pill, ["FINCA", f"  {d['finca'].replace('Finca ', '')}  ·  {len(lotes)} lotes  ·  {es(ha)} ha"])
        mover(pill, 0.45, 0.75 + alto_mapa + 0.15, 6.3, 0.42)
        texto(forma(sl, "TextBox 26"), f"{d['estado']} · Imagen satelital con delimitación de lotes "
                                       f"(Anexo B del convenio · Datum WGS 1984, UTM 19N)")
        mover(forma(sl, "TextBox 26"), y=0.75 + alto_mapa + 0.65, w=6.3)
    else:
        lat, lng = d.get("lat"), d.get("lng")
        if lat and lng:
            reemplazar_imagen(sl, pic, satelital(lat, lng))
        runs(pill, ["FINCA", f"  {d['finca'].replace('Finca ', '')}  ·  {len(lotes)} lotes"])
        mover(pill, w=4.2)
        texto(forma(sl, "TextBox 26"),
              f"{d['estado']} · {d['municipio']} · Imagen satelital de la finca (Esri World Imagery)")
        # El hueco del segundo mapa pasa a ser la tabla de lotes.
        clonar(sl, forma(sl, "Rounded Rectangle 56"), 0.45, 3.96, 6.0, 3.01)
        cols = [("LOTE", 0.7, 2.0, "l"), ("HA", 2.75, 0.9, "r"), ("SIEMBRA", 3.85, 0.9, "ctr"),
                ("pH", 4.85, 0.6, "ctr"), ("CAL", 5.5, 0.75, "ctr")]
        for t, x, w, al in cols:
            texto(clonar(sl, cab, x, 4.1, w, 0.25), t, sz=9, align=al)
        paso = min(0.27, 2.45 / max(len(filas), 1))
        for i, l in enumerate(filas):
            y = 4.42 + i * paso
            s_ = suelo.get(l["lote_id"])
            vals = [l["nombre"], es(l["ha"]), ddmm(l["fecha_siembra"]) if l["fecha_siembra"] else "—",
                    es(s_["ph"]) if s_ else "s/d", "No" if s_ and not s_["dosis_cal_kg_ha"] else ("Sí" if s_ else "s/d")]
            for (t, x, w, al), v in zip(cols, vals):
                texto(clonar(sl, cuerpo, x, y, w, paso), v, sz=10.5, bold=(t == "LOTE"), align=al,
                      color=VERDE_OSC if t == "LOTE" else "475569")

    runs(forma(sl, "TextBox 31"), [str(len(lotes))])
    runs(forma(sl, "TextBox 34"), [es(ha, 2 if len(lotes) > 3 else 1)])
    t_cal = sum((suelo.get(l["lote_id"]) or {}).get("dosis_cal_kg_ha", 0) * l["ha"] for l in lotes) / 1000
    runs(forma(sl, "TextBox 37"), [f"{es(t_cal, 0 if t_cal == 0 else 1)} t"])
    texto(forma(sl, "TextBox 38"), "de cal recomendada")

    # pH: una marca fina por lote y una sola etiqueta con el rango.
    phs = [s["ph"] for s in suelo.values()]
    xph = lambda ph: 7.68 + (min(max(ph, 4.6), 8.3) - 5) * 1.463
    linea = forma(sl, "Connector 53")
    if phs:
        lo, hi = min(phs), max(phs)
        for ph in phs:
            clonar(sl, linea, xph(ph))
        pill = forma(sl, "Rounded Rectangle 52")
        texto(pill, f"{len(phs)} lotes · pH {es(lo)} – {es(hi)}")
        mover(pill, x=min(max((xph(lo) + xph(hi)) / 2 - 1.1, 9.75), 10.6), w=2.2)
    borrar(linea, forma(sl, "Connector 55"), forma(sl, "Rounded Rectangle 54"))

    def clase(ph):
        ph = round(ph, 1)
        return ("ácido" if ph < 5.5 else "ligeramente ácido" if ph < 6.5 else "neutro" if ph < 7.3
                else "ligeramente alcalino")
    con_cal = [l["nombre"] for l in lotes if (suelo.get(l["lote_id"]) or {}).get("dosis_cal_kg_ha")]
    if phs:
        rango = f"pH de {es(min(phs), 1)} a {es(max(phs), 1)}"
        clases = sorted({clase(p) for p in phs}, key=["ácido", "ligeramente ácido", "neutro", "ligeramente alcalino"].index)
        msg = f"Tus lotes están en rango {' a '.join(clases)} ({rango}). "
        if con_cal:
            msg += f"El laboratorio recomendó cal en {', '.join(con_cal)} para corregir la acidez antes de nutrir."
        else:
            msg += ("El laboratorio no recomendó cal en ningún lote: el pH no limita la disponibilidad de "
                    "nutrientes, así que nos enfocamos en una nutrición a la medida.")
    else:
        msg = f"{PENDIENTE}: no hay análisis de suelo cargado para este agricultor."
    texto(cuerpo, msg)
    mover(cuerpo, y=5.8, h=1.0)


def _badge(sl, base, x, y, valor, lo, hi):
    estado, color = ("BAJO", DORADO) if valor < lo else ("ALTO", OXIDO) if valor > hi else ("IDEAL", VERDE)
    b = clonar(sl, base, x, y, 0.62, 0.2)
    relleno(b, color)
    texto(b, estado, sz=8)
    return estado


def lamina_suelo_resultados(sl, d):
    """Por id: en esta lámina la plantilla repite nombres (TextBox 75..81) entre tarjetas y prescripción."""
    I = lambda i: por_id(sl, i)
    lotes = [l for l in sorted(d["lotes"], key=lambda l: -l["ha"]) if l["lote_id"] in d["suelo"]]
    card, cab, valor, badge, lote_lbl = I(75), I(77), I(80), I(81), I(79)
    alto_tabla = 0.55 + 0.25 * max(len(lotes), 1)
    mover(card, 0.45, 1.78, 8.1, alto_tabla)
    cols = [("LOTE", 0.7, 1.6), ("pH", 2.4, 0.6), ("Ca / Mg  ·  ideal 3 – 5", 3.15, 2.4),
            ("Ca / K  ·  ideal 10 – 20", 5.75, 2.6)]
    for t, x, w in cols:
        texto(clonar(sl, cab, x, 1.9, w, 0.2), t, sz=9)
    conteo = {"Ca/Mg": {}, "Ca/K": {}}
    for i, l in enumerate(lotes):
        s, y = d["suelo"][l["lote_id"]], 2.2 + i * 0.25
        texto(clonar(sl, lote_lbl, 0.7, y + 0.03, 1.7, 0.2), l["nombre"], sz=10, bold=True, color=VERDE_OSC, align="l")
        texto(clonar(sl, valor, 2.4, y, 0.6, 0.22), es(s["ph"]), sz=11, align="l")
        texto(clonar(sl, valor, 3.15, y, 0.7, 0.22), es(s["ca_mg"], 1), sz=11, align="l")
        e1 = _badge(sl, badge, 3.95, y + 0.02, s["ca_mg"], 3, 5)
        texto(clonar(sl, valor, 5.75, y, 0.7, 0.22), es(s["ca_k"], 1), sz=11, align="l")
        e2 = _badge(sl, badge, 6.55, y + 0.02, s["ca_k"], 10, 20)
        conteo["Ca/Mg"][e1] = conteo["Ca/Mg"].get(e1, 0) + 1
        conteo["Ca/K"][e2] = conteo["Ca/K"].get(e2, 0) + 1
    # Las tarjetas de dos lotes de la plantilla ya no aplican.
    borrar(*[I(i) for i in range(76, 95)])

    # ¿Qué significa?
    q_card, q_tit, q_txt = I(95), I(96), I(97)
    mover(q_card, 8.75, 1.78, 4.05, alto_tabla)
    mover(q_tit, 8.97, 1.92)
    mover(q_txt, 8.97, 2.25, 3.65, alto_tabla - 0.6)
    n = len(lotes)
    cm, ck = conteo["Ca/Mg"], conteo["Ca/K"]
    partes = []
    if n and cm.get("IDEAL", 0) == n:
        partes.append("El equilibrio calcio/magnesio está en el rango ideal en todos tus lotes.")
    elif n:
        partes.append(f"Calcio/magnesio: {cm.get('IDEAL', 0)} de {n} lotes en rango ideal"
                      + (f", {cm['ALTO']} alto" if cm.get("ALTO") else "")
                      + (f", {cm['BAJO']} bajo" if cm.get("BAJO") else "") + ".")
    if ck.get("BAJO"):
        partes.append(f"En {ck['BAJO']} de {n} lotes la relación calcio/potasio está por debajo del ideal: "
                      "el potasio es abundante frente al calcio, así que el plan no necesita potasio extra.")
    if ck.get("ALTO"):
        partes.append(f"En {ck['ALTO']} de {n} lotes el calcio domina sobre el potasio: lo equilibramos con la nutrición.")
    texto(q_txt, " ".join(partes) if n else f"{PENDIENTE}: sin análisis de suelo cargado.", sz=11)

    # Prescripción: la caja baja para dejar sitio a la tabla.
    dy = Inches((1.78 + alto_tabla + 0.25) - 4.25)
    for i in range(52, 65):
        I(i).top = I(i).top + dy
    pres = d["prescripcion"]
    ha = {l["nombre"]: l["ha"] for l in d["lotes"]}
    slots = [(58, 59, 60), (61, 62, 63)]
    if pres:
        p = pres[0]
        lotes_p = list(p["por_lote"].items())
        texto(I(54), [[(p["producto"], {})], [("enmienda aplicada según el análisis", {"sz": 11})]])
        dosis = p["dosis_sacos_ha"]
        runs(I(56), [f"{es(dosis, 1 if dosis % 1 else 0)} {'saco' if dosis == 1 else 'sacos'} / ha"])
        texto(I(57), "en ambos lotes" if len(lotes_p) == len(d["lotes"]) == 2 else
              f"en {len(lotes_p)} de {len(d['lotes'])} lotes" if len(lotes_p) > 1 else f"solo en {lotes_p[0][0]}")
        if len(lotes_p) > 2:
            lotes_p = [("TOTAL", sum(v for _, v in lotes_p))]
        for (a, b, c), (nom, sacos) in zip(slots, lotes_p):
            texto(I(a), f"LOTE {nom.upper()}" if nom != "TOTAL" else "TOTAL APLICADO")
            runs(I(b), [es(sacos), " sacos"])
            texto(I(c), f"total del lote · {es(ha.get(nom, 0))} ha" if nom != "TOTAL" else "suma de los lotes")
        for a, b, c in slots[len(lotes_p):]:
            borrar(I(a), I(b), I(c))
        otros = [x for x in ("Cal", "KMAG", "MagniPlus") if x not in {q["producto"] for q in pres}]
        texto(I(64), [f"{x}: 0" for x in otros] + ["no requeridos"])
    else:
        texto(I(54), "Sin enmiendas")
        runs(I(56), ["0 sacos / ha"])
        texto(I(57), "el suelo no las requirió")
        for trio in slots:
            borrar(*[I(i) for i in trio])
        texto(I(64), ["Cal: 0", "KMAG: 0", "MagniPlus: 0"])


LABORES = ["rastra", "enmienda", "siembra", "controles", "reabono1", "reabono2", "cosecha"]


def lamina_labores(sl, d):
    lotes = sorted(d["lotes"], key=lambda l: l["nombre"])
    ha_plan = sum(l["ha"] for l in lotes)
    runs(forma(sl, "TextBox 24"), [d["agricultor"]])
    runs(forma(sl, "TextBox 25"), [es(ha_plan)])
    runs(forma(sl, "TextBox 27"), [f"planificadas a sembrar en {len(lotes)} lotes"])

    aplica = {k: any(l["labores"][k] for l in lotes) for k in LABORES}
    hecho = sum(l["labores"][k] for l in lotes for k in LABORES if aplica[k])
    total = sum(l["ha"] for l in lotes for k in LABORES if aplica[k])
    avance = hecho / total * 100 if total else 0
    texto(forma(sl, "Rounded Rectangle 28"),
          "100% DE AVANCE EN CADA LABOR" if avance >= 99.95 else f"{avance:.0f}% DE AVANCE EN LAS LABORES")

    base_cell, base_na = forma(sl, "Rounded Rectangle 41"), forma(sl, "Rounded Rectangle 44")
    base_val, base_na_txt = forma(sl, "TextBox 42"), forma(sl, "TextBox 46")
    base_nom = forma(sl, "TextBox 39")
    xs = [8.05, 8.73, 9.41, 10.09, 10.77, 11.45, 12.13]
    paso = min(0.33, 3.0 / (len(lotes) + 1))
    y0 = 2.1
    pendientes = []

    def celda(x, y, v, ref, na=False):
        if na:
            c = clonar(sl, base_na, x, y, 0.62, paso - 0.04)
            texto(c, "—", like=base_na_txt, sz=9)
            return
        c = clonar(sl, base_cell, x, y, 0.62, paso - 0.04)
        completo = v >= ref * 0.99
        if not completo:
            relleno(c, DORADO)
        texto(c, es(v), like=base_val, sz=9.5)

    for i, l in enumerate(lotes):
        y = y0 + i * paso
        texto(clonar(sl, base_nom, 6.85, y + 0.02, 1.2, paso - 0.04), l["nombre"], sz=10.5)
        for x, k in zip(xs, LABORES):
            v = l["labores"][k]
            celda(x, y, v, l["ha"], na=not aplica[k])
            if aplica[k] and v < l["ha"] * 0.99:
                pendientes.append((l["nombre"], k, v))
    linea = forma(sl, "Connector 86")
    yt = y0 + len(lotes) * paso + 0.04
    mover(linea, y=yt)
    texto(forma(sl, "TextBox 87"), "TOTAL", sz=14)
    mover(forma(sl, "TextBox 87"), y=yt + 0.06, h=0.3)
    texto(forma(sl, "TextBox 88"), f"{es(ha_plan)} ha plan", sz=9)
    mover(forma(sl, "TextBox 88"), y=yt + 0.33, h=0.2)
    for x, k in zip(xs, LABORES):
        celda(x, yt + 0.1, sum(l["labores"][k] for l in lotes), ha_plan, na=not aplica[k])

    # Fuera las filas de la plantilla.
    plantilla = [f"TextBox {i}" for i in list(range(38, 41)) + list(range(42, 70)) + list(range(89, 110))] + \
                [f"Rounded Rectangle {i}" for i in list(range(41, 84, 3)) + list(range(89, 108, 3))]
    for nombre in plantilla:
        try:
            s = forma(sl, nombre)
        except IndexError:
            continue
        if s.top < Inches(5.2):
            borrar(s)

    ley_y = yt + 0.62
    for nombre, dy in [("Rounded Rectangle 110", 0.03), ("TextBox 111", 0), ("Rounded Rectangle 112", 0.03),
                       ("TextBox 113", 0.05)]:
        mover(forma(sl, nombre), y=ley_y + dy)
    if pendientes:
        p = clonar(sl, forma(sl, "Rounded Rectangle 110"), 11.15, ley_y + 0.03)
        relleno(p, DORADO)
        mover(forma(sl, "TextBox 113"), w=1.6)
        texto(clonar(sl, forma(sl, "TextBox 113"), 11.45, ley_y + 0.05, 1.4), "Sin registro completo")
    caja, txt = forma(sl, "Rounded Rectangle 114"), forma(sl, "TextBox 115")
    mover(caja, y=ley_y + 0.4, h=7.2 - (ley_y + 0.4))
    mover(txt, y=ley_y + 0.48, h=7.2 - (ley_y + 0.62))
    nombres_lab = {"rastra": "rastra", "siembra": "siembra", "controles": "controles", "reabono1": "reabono 1",
                   "reabono2": "reabono 2", "cosecha": "cosecha", "enmienda": "enmienda"}
    completos = len({l["nombre"] for l in lotes} - {p[0] for p in pendientes})
    n_aplica = sum(aplica.values())
    msg = (f"{'Todos tus lotes' if completos == len(lotes) else f'{completos} de {len(lotes)} lotes'} "
           f"completaron las {n_aplica} labores aplicadas en el 100% de su superficie.")
    for nom, k, v in pendientes:
        msg += f" {nom}: {nombres_lab[k]} {'sin registro' if v == 0 else f'registrada en {es(v)} ha'}."
    if not aplica["enmienda"]:
        msg += " La enmienda no se aplicó en este ciclo."
    texto(txt, msg, sz=11 if len(msg) > 200 else 12)


def lamina_incidencias(sl, d):
    inc = d["incidencias"][:14]
    n = len(inc)
    colores = {"plagas": VERDE_CLARO, "maleza": VERDE_CLARO, "mazorca": VERDE_CLARO}
    oval, fecha, fase, etq = (forma(sl, "Oval 12"), forma(sl, "TextBox 13"), forma(sl, "TextBox 14"),
                             forma(sl, "TextBox 15"))
    eje = forma(sl, "Rectangle 10")
    sep = forma(sl, "Rectangle 8")
    mes_lbl = forma(sl, "TextBox 5")
    x0, x1, y = 1.0, 12.75, 2.8
    mover(eje, x0, y + 0.22, x1 - x0)
    paso = (x1 - x0) / max(n, 1)
    meses = []
    for i, it in enumerate(inc):
        cx = x0 + paso * (i + 0.5)
        mes = date.fromisoformat(it["fecha"]).month
        if not meses or meses[-1][0] != mes:
            meses.append([mes, cx, cx])
        meses[-1][2] = cx
        o = clonar(sl, oval, cx - 0.21, y, 0.42, 0.42)
        relleno(o, colores[it["reto"]])
        texto(o, str(i + 1))
        texto(clonar(sl, fecha, cx - 0.45, y - 0.32, 0.9), ddmm(it["fecha"]))
        sub = it["fase"] or (f"{it['dds']} DDS" if it["dds"] is not None else "")
        texto(clonar(sl, fase, cx - 0.45, y + 0.5, 0.9), sub)
        texto(clonar(sl, etq, cx - 0.42, y + 0.72, 0.84, 0.55), [[(it["lote"], {"b": True})], [(it["etiqueta"], {})]],
              sz=8.5)
    for j, (mes, a, b) in enumerate(meses):
        texto(clonar(sl, mes_lbl, (a + b) / 2 - 0.7, 1.7, 1.4), MESES[mes], align="ctr")
        if j:
            xs = (meses[j - 1][2] + a) / 2
            clonar(sl, sep, xs, 1.7, None, 2.4)
    for nombre in ["TextBox 5", "TextBox 6", "TextBox 7", "Rectangle 8", "Rectangle 9", "Rectangle 28",
                   "TextBox 11", "TextBox 29"] + \
            [f"Oval {i}" for i in (12, 16, 20, 24, 30, 34, 38, 42)] + \
            [f"TextBox {i}" for i in list(range(13, 28)) + list(range(31, 46)) if i not in (16, 20, 24)] + \
            [f"TextBox {i}" for i in (15, 19, 23, 27)]:
        try:
            borrar(forma(sl, nombre))
        except IndexError:
            pass
    texto(forma(sl, "TextBox 4"), "TU TEMPORADA DE INCIDENCIAS · LÍNEA DE TIEMPO")

    # Leyenda y reto por categoría
    # Mismo tag que la plantilla de Lorena: todas las incidencias están revisadas/atendidas,
    # así que van todas del mismo color y la leyenda es una sola. La categoría la dan los retos.
    texto(forma(sl, "TextBox 47"), f"{ESTADO_INCIDENCIA} ({n})")
    x = 2.42
    por_mes = {}
    for it in inc:
        m = date.fromisoformat(it["fecha"]).month
        por_mes[m] = por_mes.get(m, 0) + 1
    if por_mes:
        mm = max(por_mes, key=por_mes.get)
        texto(forma(sl, "TextBox 48"), f"{MESES[mm].capitalize()} concentró {por_mes[mm]} de {n} reportes  ·  "
                                       "DDS: días tras la siembra")
        mover(forma(sl, "TextBox 48"), x=x + 0.1, w=8.5 - x)
    retos = [("plagas", "PLAGAS EN PLÁNTULA"), ("maleza", "MALEZAS"), ("mazorca", "MAZORCA AL CIERRE")]
    tarjetas = [("TextBox 51", "TextBox 52", "TextBox 53"), ("TextBox 55", "TextBox 56", "TextBox 57"),
                ("TextBox 59", "TextBox 60", "TextBox 61")]
    for j, ((clave, titulo), (a, b, c)) in enumerate(zip(retos, tarjetas)):
        items = [it for it in inc if it["reto"] == clave]
        tit = forma(sl, a)
        texto(tit, f"RETO {j + 1} · {titulo}")
        texto(forma(sl, b), f"{len(items)} de {n}")
        if items:
            dds = [it["dds"] for it in items if it["dds"] is not None]
            lotes = sorted({it["lote"] for it in items})
            rango = f"entre {min(dds)} y {max(dds)} DDS" if dds and min(dds) != max(dds) else (f"a los {dds[0]} DDS" if dds else "")
            desc = {"plagas": "reportes de gusano cogollero y larvas",
                    "maleza": "controles dirigidos de malezas",
                    "mazorca": "mazorcas con agobio o daño"}[clave]
            lugar = lotes[0] if len(lotes) == 1 else f"{len(lotes)} lotes"
            texto(forma(sl, c), f"{desc} {rango}, en {lugar}")
        else:
            texto(forma(sl, c), "sin reportes en este ciclo")


def lamina_clima(sl, d):
    png, parciales = grafico_lluvia(d)
    pic = forma(sl, "Imagen 2")
    reemplazar_imagen(sl, pic, png)
    mover(pic, 0.31, 1.78, 12.73, 3.35)
    for nombre in ["Rectángulo 12", "Conector recto 12", "Rectángulo 17", "Imagen 14"]:
        borrar(forma(sl, nombre))
    c = d["clima"]
    primer = d["agricultor"].split()[0]
    texto(forma(sl, "TextBox 24"), f"Fuente: estación meteorológica {c['estacion'] or PENDIENTE}, "
                                   f"la estación Saturno asignada a la finca de {primer}.")
    nota = "Abril parcial: la estación empezó a registrar ese mes. " if 4 in parciales else ""
    # KPIs nativos en lugar de la captura de la plantilla.
    flor = c["floracion"]
    kpis = [("LOTES MOSTRADOS", str(len(d["lotes"])), "con la misma estación"),
            ("LLUVIA EN EL CICLO", f"{es(c['lluvia_ciclo_mm'], 1)} mm" if c["lluvia_ciclo_mm"] else PENDIENTE,
             "desde la siembra a hoy"),
            ("LLUVIA EN FLORACIÓN", f"{es(min(f['mm'] for f in flor), 0)} – {es(max(f['mm'] for f in flor), 0)} mm"
             if flor else PENDIENTE, f"meta {es(flor[0]['meta_mm'], 0)} mm en 30 días" if flor else ""),
            ("RACHA SECA MÁXIMA", f"{max(f['racha_max_dias_secos'] for f in flor)} días" if flor else PENDIENTE,
             "en la ventana de floración")]
    x = 1.2
    for tit, val, sub in kpis:
        t = tarjeta(sl, x, 6.1, 2.6, 0.95)
        tb = sl.shapes.add_textbox(Inches(x + 0.15), Inches(6.15), Inches(2.35), Inches(0.85))
        tf = tb.text_frame
        tf.word_wrap = True
        for k, (txt, sz, b, col) in enumerate([(tit, 8, True, GRIS), (val, 17, True, VERDE_OSC), (sub, 8, False, GRIS)]):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            r = p.add_run()
            r.text = txt
            r.font.size, r.font.bold, r.font.name = Pt(sz), b, "72 Condensed"
            r.font.color.rgb = RGBColor.from_string(col)
        x += 2.8
    if nota:
        tb = sl.shapes.add_textbox(Inches(0.31), Inches(5.15), Inches(6), Inches(0.25))
        r = tb.text_frame.paragraphs[0].add_run()
        r.text = nota + "Puntos huecos = mes con datos incompletos."
        r.font.size, r.font.name = Pt(9), "72 Condensed"
        r.font.color.rgb = RGBColor.from_string(GRIS)


def lamina_mecanizacion(sl, d):
    pagos = d.get("mecanizacion_pagos") or {}
    aviso = "" if pagos else f"  ·  {PENDIENTE}: cargar la liquidación de APC"
    texto(forma(sl, "TextBox 104"), f"{d['agricultor']} · montos en USD{aviso}")
    campos = {"TextBox 107": "real_inicial", "TextBox 111": "pago1", "TextBox 115": "pago2",
              "TextBox 119": "pago3", "TextBox 127": "restante", "TextBox 131": "ajustes",
              "TextBox 135": "neto", "TextBox 139": "indexado", "TextBox 143": "comisiones",
              "TextBox 147": "total", "TextBox 151": "aj_encaladora", "TextBox 155": "aj_dron",
              "TextBox 159": "aj_compras"}
    for nombre, k in campos.items():
        v = pagos.get(k)
        runs(forma(sl, nombre), [es(v) if v is not None else "—"])


def lamina_pl(sl, d):
    pl = d["pl"]
    runs(forma(sl, "TextBox 7"), [d["agricultor"], f"{d['municipio']}, {d['estado']}"])
    rs = forma(sl, "TextBox 7").text_frame._txBody.xpath(".//a:r")
    if len(rs) < 2:
        texto(forma(sl, "TextBox 7"), [d["agricultor"], f"{d['municipio']}, {d['estado']}"])
    runs(forma(sl, "TextBox 10"), [es(pl["ha_cosechadas"], 1), " ha"])
    texto(forma(sl, "TextBox 11"), f"{len(d['lotes'])} lotes · {es(pl['ha_plan'], 1)} ha sembradas")
    runs(forma(sl, "TextBox 14"), [es(pl["rend_t_ha"]), " ton/ha"])
    texto(forma(sl, "TextBox 15"), f"{pl['rend_t_ha'] / pl['meta_t_ha'] * 100:.0f}% del plan "
                                   f"({es(pl['meta_t_ha'], 0)} ton/ha)")
    ton = forma(sl, "TextBox 18")
    runs(ton, [es(pl["toneladas"], 1), " ton"])
    cifra = ton.text_frame.paragraphs[0].runs[0]
    if len(cifra.text) > 5 and cifra.font.size:  # "2.269,2" no cabe al tamaño de "322,7"
        cifra.font.size = Pt(cifra.font.size.pt * 5.5 / len(cifra.text))
    texto(forma(sl, "TextBox 19"), f"de {es(pl['ton_plan'], 0)} ton planificadas")


def lamina_flujo(sl, d):
    fc = d.get("flujo_caja") or {}
    aviso = "" if fc else f"  ·  {PENDIENTE}: cargar el flujo de caja"
    texto(forma(sl, "TextBox 104"), f"{d['agricultor']} · montos en USD{aviso}")
    tabla = forma(sl, "Tabla 8").table
    for fila in tabla.rows:
        concepto = fila.cells[0].text.strip()
        if not concepto:
            continue
        for j, cel in enumerate(list(fila.cells)[1:3]):
            v = (fc.get(concepto) or [None, None])[j]
            tf = cel.text_frame
            r = tf.paragraphs[0].runs
            if r:
                r[0].text = es(v) if v is not None else "—"
                for extra in r[1:]:
                    extra.text = ""


def notas_id(prs, d):
    """La nota de la lámina de laboratorio citaba un lote de Lorena."""
    sl = prs.slides[10]
    if not sl.has_notes_slide:
        return
    tf = sl.notes_slide.notes_text_frame
    plaga = next((i for i in d["incidencias"] if i["reto"] == "plagas" and i["origen"] == "actividad"), None)
    for p in tf.paragraphs:
        for r in p.runs:
            if "Vincular con" in r.text:
                r.text = r.text.split("Vincular con")[0]
    if plaga:
        tf.add_paragraph().text = (f"Vincular con {plaga['lote']}: el {ddmm(plaga['fecha'])} se reportó "
                                   f"«{plaga['texto'].strip()}».")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("datos")
    ap.add_argument("salida")
    ap.add_argument("--plantilla", default=PLANTILLA)
    ap.add_argument("--nombre-corto")
    a = ap.parse_args()
    d = json.load(open(a.datos, encoding="utf-8"))
    prs = Presentation(a.plantilla)
    s = prs.slides
    lamina_portada_agricultor(s[1], d, a.nombre_corto or nombre_corto(d["agricultor"]))
    lamina_suelo_mapa(s[2], d)
    lamina_suelo_resultados(s[3], d)
    lamina_labores(s[5], d)
    lamina_incidencias(s[6], d)
    lamina_clima(s[7], d)
    lamina_mecanizacion(s[13], d)
    lamina_pl(s[14], d)
    lamina_flujo(s[15], d)
    notas_id(prs, d)
    prs.save(a.salida)


if __name__ == "__main__":
    main()
